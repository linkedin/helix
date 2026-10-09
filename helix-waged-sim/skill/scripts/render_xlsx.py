#!/usr/bin/env python3
# Licensed to the Apache Software Foundation (ASF) under one or more contributor license agreements.
# See the NOTICE file distributed with this work for additional information regarding copyright
# ownership. The ASF licenses this file to you under the Apache License, Version 2.0.
"""Render a colour-coded report.xlsx from a waged-sim run folder (run.json, rounds.jsonl, nodes CSVs).

Usage: render_xlsx.py <run folder> [output.xlsx]

Needs openpyxl. If it is missing:
  python3 -m venv ~/.waged-sim/venv && ~/.waged-sim/venv/bin/pip install openpyxl
  ~/.waged-sim/venv/bin/python render_xlsx.py <run folder>
"""
import csv
import json
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit(__doc__)

GREEN, YELLOW, RED = (198, 239, 206), (255, 235, 156), (255, 199, 206)
HEADER = PatternFill("solid", fgColor="1F3864")
VERDICT = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "ERROR": "FFEB9C"}


def mix(a, b, t):
    return tuple(round(x + (y - x) * t) for x, y in zip(a, b))


def scale(value, low, mid, high):
    if value <= low:
        rgb = GREEN
    elif value >= high:
        rgb = RED
    elif value <= mid:
        rgb = mix(GREEN, YELLOW, (value - low) / (mid - low))
    else:
        rgb = mix(YELLOW, RED, (value - mid) / (high - mid))
    return PatternFill("solid", fgColor="%02X%02X%02X" % rgb)


def fill_for(stat, value):
    if not isinstance(value, (int, float)):
        return None
    if stat.startswith("skew."):
        return scale(value, 1.0, 1.15, 1.4)
    if stat.startswith(("maxUtil.", "util.")) or "UtilPct" in stat:
        return scale(value, 40, 70, 100)
    return None


def write_table(ws, header, rows, fills=None):
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
        if fills:
            for index, (name, value) in enumerate(zip(header, row), start=1):
                fill = fills(name, value)
                if fill is not None:
                    ws.cell(row=ws.max_row, column=index).fill = fill
    for index, name in enumerate(header, start=1):
        width = max([len(str(name))] + [len(str(r[index - 1])) for r in rows if index - 1 < len(r)])
        ws.column_dimensions[get_column_letter(index)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "B2"


def number(value):
    if isinstance(value, str):
        try:
            return float(value) if "." in value else int(value)
        except ValueError:
            return value
    return value


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    run_dir = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(run_dir, "report.xlsx")
    summary = json.load(open(os.path.join(run_dir, "run.json")))
    rounds = [json.loads(line) for line in open(os.path.join(run_dir, "rounds.jsonl")) if line.strip()]
    focus = summary.get("focusKey")
    scenario = summary.get("scenario", {})
    stats = list(scenario.get("report", {}).get("stats") or [])
    if not stats:
        stats = [f"skew.top.{focus}", f"skew.all.{focus}", f"maxUtil.top.{focus}", "skew.topCount",
                 "moves.replicas", "moves.topState", "missingTopState", "underReplicated"]
    for stat in scenario.get("exit", {}).get("conditionStats", []):
        if stat not in stats:
            stats.append(stat)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    variants = summary.get("variants", [])
    write_table(ws, ["Variant", "Result", "Round", "Reason", "Elapsed"],
                [[v["name"], v["status"], v["round"], v["reason"], v.get("elapsed", "")] for v in variants])
    for row in range(2, ws.max_row + 1):
        status = ws.cell(row=row, column=2).value
        if status in VERDICT:
            ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=VERDICT[status])
    cluster = summary.get("cluster", {})
    ws.append([])
    for label, value in [("Scenario", scenario.get("name")), ("Cluster", cluster.get("cluster")),
                         ("Source", cluster.get("source")), ("Captured", cluster.get("capturedAt")),
                         ("Mode", scenario.get("mode")), ("Focus key", focus),
                         ("Fidelity", "; ".join(cluster.get("fidelity", []))),
                         ("Command", summary.get("command"))]:
        ws.append([label, value])

    ws = wb.create_sheet("Rounds")
    header = ["Variant", "Round", "Events", "Passes"] + stats
    rows = []
    for record in rounds:
        passes = "+".join(k for k, v in record.get("passes", {}).items()
                          if isinstance(v, (int, float)) and v > 0 and k not in ("emergency", "failures"))
        events = "; ".join(record.get("events", []))
        rows.append([record["variant"], record["round"], events[:200], passes]
                    + [record.get("stats", {}).get(s) for s in stats])
    write_table(ws, header, rows, fill_for)

    ws = wb.create_sheet("Final")
    names = []
    for v in variants:
        for key in list(v.get("start", {})) + list(v.get("end", {})):
            if key not in names:
                names.append(key)
    header = ["Stat"] + [f"{v['name']} start" for v in variants] + [f"{v['name']} end" for v in variants]
    rows = [[n] + [v.get("start", {}).get(n) for v in variants] + [v.get("end", {}).get(n) for v in variants]
            for n in names]
    write_table(ws, header, rows, lambda name, value: None)
    for row in ws.iter_rows(min_row=2):
        stat = row[0].value
        for cell in row[1:]:
            fill = fill_for(stat, cell.value)
            if fill is not None:
                cell.fill = fill

    for v in variants[:5]:
        path = os.path.join(run_dir, f"nodes_{v['name']}_{v['rounds']}.csv".replace("/", "_"))
        if not os.path.exists(path):
            continue
        with open(path) as handle:
            reader = list(csv.DictReader(handle))
        key = f"topUtilPct.{focus}"
        reader.sort(key=lambda r: -float(r.get(key) or 0))
        columns = ["instance", "zone", "serving", key, f"allUtilPct.{focus}", "topCount", "replicaCount",
                   "topResources"]
        ws = wb.create_sheet(f"Nodes {v['name']}"[:31])
        write_table(ws, columns, [[number(r.get(c)) for c in columns] for r in reader], fill_for)
    wb.save(out)
    print(f"Report: {out}")


if __name__ == "__main__":
    main()
